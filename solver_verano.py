from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from calendar import monthrange

# ==================================================
# CONFIGURACION
# ==================================================

OPERADORES = list("ABCDEFGHIJKL")

INICIO = date(2027, 1, 1)
FIN = date(2027, 12, 31)

COLORES = {
    "M": "92D050",
    "T": "FFC000",
    "N": "7030A0",
    "4": "D9D9D9"
}

DIAS_SEMANA = [
    "L",
    "M",
    "X",
    "J",
    "V",
    "S",
    "D"
]

MESES = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE"
]

# ==================================================
# FUNCIONES CICLO
# ==================================================

def es_verano(fecha):
    return fecha.month in [6, 7, 8, 9]


def longitud_ciclo(fecha):
    if es_verano(fecha):
        return 9
    return 12


def turno_desde_posicion(pos, fecha):

    if pos in [0, 1\]:
        return "M"

    if pos in [2, 3\]:
        return "T"

    if pos in [4, 5\]:
        return "N"

    return "4"


# ==================================================
# DESFASES INICIALES
# (los que encontró OR-ls)
# ==================================================

OFFSETS = {
    "A": 6,
    "B": 5,
    "C": 3,
    "D": 7,
    "E": 10,
    "F": 9,
    "G": 1,
    "H": 8,
    "I": 2,
    "J": 0,
    "K": 4,
    "L": 11,
}

# ==================================================
# GENERAR ESTADOS
# ==================================================

cuadrante = {}

for op in OPERADORES:

    posicion = OFFSETS[op]

    fecha = INICIO

    while fecha <= FIN:

        longitud = longitud_ciclo(fecha)

        posicion = posicion % longitud

        turno = turno_desde_posicion(
            posicion,
            fecha
        )

        cuadrante[(op, fecha)] = turno

        posicion += 1

        fecha += timedelta(days=1)

# ==================================================
# EXCEL
# ==================================================

wb = Workbook()

ws = wb.active
ws.title = "Cuadrante"

fila = 1
indice_fecha = INICIO

for mes in range(1, 13):

    dias_mes = monthrange(2027, mes)[1]

    ws.cell(
        fila,
        1,
        f"{MESES[mes-1]} 2027"
    )

    fila += 1

    # días

    ws.cell(fila, 1, "Operador")

    for dia in range(1, dias_mes + 1):

        ws.cell(
            fila,
            dia + 1,
            f"{dia:02d}"
        )

    fila += 1

    # inicial día

    ws.cell(fila, 1, "Dia")

    for dia in range(1, dias_mes + 1):

        fecha = date(
            2027,
            mes,
            dia
        )

        ws.cell(
            fila,
            dia + 1,
            DIAS_SEMANA[
                fecha.weekday()
            ]
        )

    fila += 1

    inicio_operadores = fila

    # operadores

    for op in OPERADORES:

        ws.cell(
            fila,
            1,
            op
        )

        for dia in range(1, dias_mes + 1):

            fecha = date(
                2027,
                mes,
                dia
            )

            turno = cuadrante[
                (op, fecha)
            ]

            c = ws.cell(
                fila,
                dia + 1,
                turno
            )

            c.fill = PatternFill(
                "solid",
                fgColor=COLORES[turno]
            )

        fila += 1

    # cobertura

    for turno in ["M", "T", "N"]:

        ws.cell(
            fila,
            1,
            f"Cob {turno}"
        )

        for dia in range(1, dias_mes + 1):

            fecha = date(
                2027,
                mes,
                dia
            )

            conteo = 0

            for op in OPERADORES:

                if cuadrante[
                    (op, fecha)
                ] == turno:

                    conteo += 1

            ws.cell(
                fila,
                dia + 1,
                conteo
            )

        fila += 1

    fila += 2

# ==================================================
# GUARDAR
# ==================================================

archivo = "Cuadrante_2027_Estado_Diario.xlsx"

wb.save(archivo)

print()
print("Excel generado:")
print(archivo)