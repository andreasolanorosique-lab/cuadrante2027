from ortools.sat.python import cp_model
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import date
from calendar import monthrange

# ==================================================
# CONFIGURACION
# ==================================================

OPERADORES = list("ABCDEFGHIJKL")

PATRON_INVIERNO = [
    1, 1,   # MM
    2, 2,   # TT
    3, 3,   # NN
    0, 0, 0, 0, 0, 0   # 444444
]
PATRON_VERANO = [
1, 1, # MM
2, 2, # TT
3, 3, # NN
0, 0, 0 # 444
]

DIAS = 365
LONGITUD_INVIERNO = 12
LONGITUD_VERANO = 9

# ==================================================
# MODELO
# ==================================================

model = cp_model.CpModel()

turno = {}
posicion = {}
siguiente = {}

for op in OPERADORES:

    for d in range(DIAS):

        posicion[op, d] = model.NewIntVar(
            0,
            11,
            f"pos_{op}_{d}"
        )
        siguiente[op, d] = model.NewIntVar(
            0,
            11,
            f"sig_{op}_{d}"
        )
        model.Add(
            siguiente[op, d]
            ==
            posicion[op, d]
        )
        turno[op, d] = model.NewIntVar(
            0,
            11,
            f"{op}_{d}"
        )

        opciones = []
        
        #for k in range(LONGITUD_INVIERNO):

            #b = model.NewBoolVar(
                #f"{op}_{d}_{k}"
        #)

            #model.Add(
                
            #).OnlyEnforceIf(b)

            #model.Add(
                #posicion[op, d] == k
            #).OnlyEnforceIf(b)

            #model.Add(
                #posicion[op, d] != k
            #).OnlyEnforceIf(b.Not())
            
            
            #model.Add(
                #turno[op, d]
                #==
                #posicion[op, d]
            #).OnlyEnforceIf(b)

            #opciones.append(b)

        #model.AddExactlyOne(opciones)
        model.Add(
            turno[op, d]
            ==
            posicion[op, d]
        )
        if d > 0:

            model.Add(
            siguiente[op, d - 1]
            ==
            posicion[op, d]
        )
            
# ==================================================
# COBERTURA
# ==================================================

for d in range(DIAS):

    m = []
    t = []
    n = []

    for op in OPERADORES:

        es_m = model.NewBoolVar(f"m_{op}_{d}")
        es_t = model.NewBoolVar(f"t_{op}_{d}")
        es_n = model.NewBoolVar(f"n_{op}_{d}")

        model.Add(turno[op, d] == 1).OnlyEnforceIf(es_m)
        model.Add(turno[op, d] != 1).OnlyEnforceIf(es_m.Not())

        model.Add(turno[op, d] == 2).OnlyEnforceIf(es_t)
        model.Add(turno[op, d] != 2).OnlyEnforceIf(es_t.Not())

        model.Add(turno[op, d] == 3).OnlyEnforceIf(es_n)
        model.Add(turno[op, d] != 3).OnlyEnforceIf(es_n.Not())

        m.append(es_m)
        t.append(es_t)
        n.append(es_n)

    model.Add(sum(m) >= 2)
    model.Add(sum(t) >= 2)
    model.Add(sum(n) >= 2)

# ==================================================
# RESOLVER
# ==================================================

solver = cp_model.CpSolver()
solver.parameters.max_time_in_seconds = 60

status = solver.Solve(model)

print("STATUS =", solver.StatusName(status))
print()
print("PRUEBA POSICIONES A")
print()
print("PRUEBA SIGUIENTE A")

for d in range(10):

    print(
        d,
        solver.Value(posicion["A", d])
    )
for d in range(10):

    print(
        d,
        solver.Value(siguiente["A", d])
    )

if status not in (
    cp_model.OPTIMAL,
    cp_model.FEASIBLE
):
    raise ValueError("No se encontró solución")

# ==================================================
# EXCEL
# ==================================================

wb = Workbook()

ws = wb.active
ws.title = "Cuadrante"

meses = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE"
]

dias_semana = ["L", "M", "X", "J", "V", "S", "D"]

colores = {
    "M": "92D050",
    "T": "FFC000",
    "N": "7030A0",
    "4": "D9D9D9"
}

fila = 1
indice_global = 0

for mes in range(1, 13):

    dias_mes = monthrange(2027, mes)[1]

    ws.cell(
        fila,
        1,
        f"{meses[mes-1]} 2027"
    )

    ws.cell(
        fila,
        1
    ).font = Font(
        bold=True,
        size=14
    )

    fila += 1

    # fila dias

    ws.cell(fila, 1, "Operador")

    for dia in range(1, dias_mes + 1):

        ws.cell(
            fila,
            dia + 1,
            f"{dia:02d}"
        )

    fila += 1

    # fila iniciales

    ws.cell(fila, 1, "Dia")

    for dia in range(1, dias_mes + 1):

        fecha = date(
            2027,
            mes,
            dia
        )

        ws.cell(
            fila,
            dia + 1,
            dias_semana[
                fecha.weekday()
            ]
        )

    fila += 1

    inicio_operadores = fila

    # operadores

    for op in OPERADORES:

        ws.cell(
            fila,
            1,
            op
        )

        for d in range(dias_mes):

            valor = solver.Value(
                turno[
                    op,
                    indice_global + d
                ]
            )

            if valor == 1:
                texto = "M"
            elif valor == 2:
                texto = "T"
            elif valor == 3:
                texto = "N"
            else:
                texto = "4"

            c = ws.cell(
                fila,
                d + 2,
                texto
            )

            c.fill = PatternFill(
                "solid",
                fgColor=colores[texto]
            )

            c.alignment = Alignment(
                horizontal="center"
            )

        fila += 1

    # cobertura

    for letra_turno in ["M", "T", "N"]:

        ws.cell(
            fila,
            1,
            f"Cob {letra_turno}"
        )

        for d in range(dias_mes):

            conteo = 0

            for r in range(
                inicio_operadores,
                inicio_operadores + len(OPERADORES)
            ):

                if (
                    ws.cell(r, d + 2).value
                    ==
                    letra_turno
                ):
                    conteo += 1

            ws.cell(
                fila,
                d + 2,
                conteo
            )

        fila += 1

    fila += 2

    indice_global += dias_mes

# ancho columnas

ws.column_dimensions["A"].width = 15

for col in range(2, 40):
    letra = ws.cell(2, col).column_letter
    ws.column_dimensions[letra].width = 4

# hoja cobertura anual

cov = wb.create_sheet(
    "Cobertura Anual"
)

cov.append(
    ["Dia", "M", "T", "N"]
)

for d in range(DIAS):

    m = 0
    t = 0
    n = 0

    for op in OPERADORES:

        valor = solver.Value(
            turno[op, d]
        )

        if valor == 1:
            m += 1
        elif valor == 2:
            t += 1
        elif valor == 3:
            n += 1

    cov.append(
        [d + 1, m, t, n]
    )

nombre = "Cuadrante_2027_Meses.xlsx"

wb.save(nombre)

print()
print(nombre)