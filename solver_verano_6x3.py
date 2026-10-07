from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill

OPERADORES = list("ABCDEFGHIJKL")

INICIO = date(2027, 1, 1)
FIN = date(2027, 12, 31)

PATRON_INVIERNO = [
    "M", "M",
    "T", "T",
    "N", "N",
    "4", "4", "4", "4", "4", "4"
]

PATRON_VERANO = [
    "M", "M",
    "T", "T",
    "N", "N",
    "4", "4", "4"
]

OFFSETS = {
    "A": 6,
    "B": 5,
    "C": 3,
    "D": 7,
    "E": 10,
    "F": 9,
    "G": 1,
    "H": 8,
    "I": 2,
    "J": 0,
    "K": 4,
    "L": 11
}

COLORES = {
    "M": "92D050",
    "T": "FFC000",
    "N": "7030A0",
    "4": "D9D9D9"
}

def es_verano(fecha):
    return fecha.month in [6, 7, 8, 9]

cuadrante = {}

for operador in OPERADORES:

    posicion = OFFSETS[operador]

    fecha = INICIO

    while fecha <= FIN:

        if es_verano(fecha):
            patron = PATRON_VERANO
        else:
            patron = PATRON_INVIERNO

        posicion_real = posicion % len(patron)

        cuadrante[(operador, fecha)] = patron[posicion_real]

        posicion += 1
        fecha += timedelta(days=1)

print()
print("COMPROBACION COBERTURA")
print("----------------------")

fecha = INICIO

while fecha <= FIN:

    m = 0
    t = 0
    n = 0

    for operador in OPERADORES:

        turno = cuadrante[(operador, fecha)]

        if turno == "M":
            m += 1

        elif turno == "T":
            t += 1

        elif turno == "N":
            n += 1

    if m < 2 or t < 2 or n < 2:

        print(
            fecha,
            "M=", m,
            "T=", t,
            "N=", n
        )

    fecha += timedelta(days=1)

wb = Workbook()

ws = wb.active
ws.title = "Verano 6x3"

ws.cell(1, 1, "Operador")

fecha = INICIO
col = 2

while fecha <= FIN:

    ws.cell(
        1,
        col,
        fecha.strftime("%d/%m")
    )

    col += 1
    fecha += timedelta(days=1)

for fila, operador in enumerate(
    OPERADORES,
    start=2
):

    ws.cell(
        fila,
        1,
        operador
    )

    fecha = INICIO
    col = 2

    while fecha <= FIN:

        turno = cuadrante[
            (operador, fecha)
        ]

        c = ws.cell(
            fila,
            col,
            turno
        )

        c.fill = PatternFill(
            "solid",
            fgColor=COLORES[turno]
        )

        col += 1
        fecha += timedelta(days=1)

archivo = "Cuadrante_2027_Verano_6x3.xlsx"

wb.save(archivo)

print()
print("Archivo generado:")
print(archivo)